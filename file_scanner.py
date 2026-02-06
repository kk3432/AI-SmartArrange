import os
from pathlib import Path
from typing import List, Tuple, Dict
import PyPDF2
from docx import Document
from PIL import Image
import openpyxl
from config import Config

class FileScanner:
    def __init__(self, source_dir: Path = None):
        self.source_dir = source_dir or Config.SOURCE_DIR
        self.supported_extensions = Config.SUPPORTED_EXTENSIONS
        
    def scan_files(self) -> List[Path]:
        """扫描目录中的所有文件"""
        files = []
        for ext in self.supported_extensions:
            files.extend(list(self.source_dir.glob(f"*{ext}")))
            files.extend(list(self.source_dir.glob(f"*{ext.upper()}")))
        
        # 去除重复（大小写可能重复）
        return list(set(files))
    
    def extract_text_content(self, file_path: Path) -> str:
        """提取文件的文本内容"""
        content = ""
        ext = file_path.suffix.lower()
        
        try:
            if ext == '.txt' or ext == '.md':
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    content = f.read(5000)  # 限制读取长度
                    
            elif ext == '.pdf':
                with open(file_path, 'rb') as f:
                    pdf_reader = PyPDF2.PdfReader(f)
                    for page in pdf_reader.pages[:3]:  # 只读取前3页
                        content += page.extract_text() + "\n"
                        if len(content) > 5000:
                            break
                            
            elif ext in ['.doc', '.docx']:
                doc = Document(file_path)
                for para in doc.paragraphs[:50]:  # 限制段落数
                    content += para.text + "\n"
                    if len(content) > 5000:
                        break
                        
            elif ext in ['.xls', '.xlsx']:
                wb = openpyxl.load_workbook(file_path, read_only=True)
                sheet = wb.active
                for row in sheet.iter_rows(max_row=100, values_only=True):
                    row_text = ' '.join(str(cell) for cell in row if cell)
                    content += row_text + "\n"
                    if len(content) > 5000:
                        break
                        
            elif ext in ['.jpg', '.jpeg', '.png', '.gif']:
                # 对于图片文件，返回文件名和基本描述
                content = f"图片文件: {file_path.name}"
                
            else:
                # 对于其他文件类型，使用文件名
                content = f"文件名: {file_path.name}"
                
        except Exception as e:
            print(f"读取文件 {file_path.name} 时出错: {e}")
            content = f"文件名: {file_path.name}"
            
        return content[:5000]  # 确保不超过5000字符
    
    def get_files_with_content(self) -> List[Tuple[Path, str]]:
        """获取文件及其内容"""
        files = self.scan_files()
        result = []
        
        print(f"找到 {len(files)} 个文件")
        
        for file_path in files:
            if file_path.is_file():
                content = self.extract_text_content(file_path)
                result.append((file_path, content))
                
        return result
